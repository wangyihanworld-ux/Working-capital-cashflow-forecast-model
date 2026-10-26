import unittest
import pandas as pd
from working_capital.evidence import profit_cash_bridge,one_factor_sensitivity
from working_capital.synthetic import build_synthetic_dataset
from working_capital.model import build_forecast

class EvidenceTests(unittest.TestCase):
    def test_hand_calculated_bridge(self):
        f=pd.DataFrame([dict(scenario='hand',month='2026-01',revenue_mn=100,cogs_mn=60,opex_mn=10,
            opening_ar_mn=20,ending_ar_mn=30,opening_inventory_mn=40,ending_inventory_mn=45,
            opening_ap_mn=15,ending_ap_mn=18,operating_cashflow_mn=18)])
        b=profit_cash_bridge(f).iloc[0]
        self.assertEqual(b.operating_profit_mn,30)
        self.assertEqual(b.indirect_ocf_mn,18)
        self.assertEqual(b.reconciliation_difference_mn,0)
    def test_all_scenarios_reconcile(self):
        d=build_synthetic_dataset()
        for _,a in d.assumptions.iterrows():
            b=profit_cash_bridge(build_forecast(d.history,a.to_dict()))
            self.assertLess(b.reconciliation_difference_mn.abs().max(),1e-8)
    def test_one_factor_does_not_mutate_and_direction_is_explained(self):
        d=build_synthetic_dataset(); a=d.assumptions.query("scenario=='基准'").iloc[0].to_dict(); old=dict(a)
        s=one_factor_sensitivity(d.history,a); self.assertEqual(a,old)
        for driver in ['dso','dio']:
            self.assertLess(s.query('driver==@driver and delta_days>0').iloc[0].change_vs_base_mn,0)
        self.assertGreater(s.query("driver=='dpo' and delta_days>0").iloc[0].change_vs_base_mn,0)
    def test_report_contains_reconciled_cash_bridge(self):
        import tempfile
        from pathlib import Path
        from openpyxl import load_workbook
        from working_capital.reporting import generate_report
        d=build_synthetic_dataset()
        with tempfile.TemporaryDirectory() as folder:
            path=generate_report(d.history,d.assumptions,Path(folder)/'report.xlsx')
            wb=load_workbook(path,data_only=True)
            sheet=wb['利润现金桥']; rows=list(sheet.values); headers=list(rows[0])
            column=headers.index('现金勾稽差异（百万元）')
            self.assertLess(max(abs(row[column]) for row in rows[1:]),1e-8)
            self.assertEqual(wb['单因素敏感性'].max_row,7)
            self.assertIn('不含税款',wb['管理摘要']['A18'].value)
            wb.close()
