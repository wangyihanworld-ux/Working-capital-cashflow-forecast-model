import tempfile,unittest
import pandas as pd
from openpyxl import load_workbook
from working_capital.synthetic import build_synthetic_dataset
from working_capital.model import build_forecast,summarize_scenarios
from working_capital.validation import DataValidationError
from working_capital.reporting import generate_demo_artifacts,REPORT_SHEETS
class WorkingCapitalTest(unittest.TestCase):
 def setUp(self): self.d=build_synthetic_dataset(); self.base=self.d.assumptions.query("scenario=='基准'").iloc[0].to_dict()
 def test_repeatable(self): pd.testing.assert_frame_equal(self.d.history,build_synthetic_dataset().history)
 def test_history_months(self): self.assertEqual(len(self.d.history),24)
 def test_forecast_months(self): self.assertEqual(len(build_forecast(self.d.history,self.base)),12)
 def test_ar_rollforward(self):
  f=build_forecast(self.d.history,self.base); self.assertLess((f.opening_ar_mn+f.revenue_mn-f.collections_mn-f.ending_ar_mn).abs().max(),1e-8)
 def test_inventory_rollforward(self):
  f=build_forecast(self.d.history,self.base); self.assertLess((f.opening_inventory_mn+f.purchases_mn-f.cogs_mn-f.ending_inventory_mn).abs().max(),1e-8)
 def test_ap_rollforward(self):
  f=build_forecast(self.d.history,self.base); self.assertLess((f.opening_ap_mn+f.purchases_mn-f.supplier_payments_mn-f.ending_ap_mn).abs().max(),1e-8)
 def test_cash_rollforward(self):
  f=build_forecast(self.d.history,self.base); self.assertLess((f.opening_cash_mn+f.operating_cashflow_mn-f.ending_cash_mn).abs().max(),1e-8)
 def test_ocf_formula(self):
  f=build_forecast(self.d.history,self.base); self.assertLess((f.collections_mn-f.supplier_payments_mn-f.opex_mn-f.operating_cashflow_mn).abs().max(),1e-8)
 def test_ccc_formula(self):
  f=build_forecast(self.d.history,self.base); self.assertLess((f.dso+f.dio-f.dpo-f.ccc).abs().max(),1e-8)
 def test_dso_matches_assumption(self): self.assertAlmostEqual(build_forecast(self.d.history,self.base).dso.mean(),48)
 def test_scenario_count(self): self.assertEqual(len(summarize_scenarios(self.d.history,self.d.assumptions)['summary']),4)
 def test_pressure_cash_below_base(self):
  s=summarize_scenarios(self.d.history,self.d.assumptions)['summary'].set_index('scenario'); self.assertLess(s.loc['压力','minimum_cash_mn'],s.loc['基准','minimum_cash_mn'])
 def test_improvement_cash_above_base(self):
  s=summarize_scenarios(self.d.history,self.d.assumptions)['summary'].set_index('scenario'); self.assertGreater(s.loc['改善','ending_cash_mn'],s.loc['基准','ending_cash_mn'])
 def test_higher_dso_lowers_cash(self):
  slow=self.base.copy(); slow['dso']=75; self.assertLess(build_forecast(self.d.history,slow).ending_cash_mn.iloc[-1],build_forecast(self.d.history,self.base).ending_cash_mn.iloc[-1])
 def test_higher_dio_lowers_cash(self):
  high=self.base.copy(); high['dio']=90; self.assertLess(build_forecast(self.d.history,high).ending_cash_mn.iloc[-1],build_forecast(self.d.history,self.base).ending_cash_mn.iloc[-1])
 def test_higher_dpo_improves_cash(self):
  long=self.base.copy(); long['dpo']=60; self.assertGreater(build_forecast(self.d.history,long).ending_cash_mn.iloc[-1],build_forecast(self.d.history,self.base).ending_cash_mn.iloc[-1])
 def test_duplicate_month_rejected(self):
  x=pd.concat([self.d.history,self.d.history.iloc[[0]]]);
  with self.assertRaisesRegex(DataValidationError,'不能重复'): build_forecast(x,self.base)
 def test_invalid_rate_rejected(self):
  x=self.base.copy(); x['cogs_rate']=1.2
  with self.assertRaisesRegex(DataValidationError,'0到1'): build_forecast(self.d.history,x)
 def test_multiple_scenarios_rejected(self):
  with self.assertRaisesRegex(DataValidationError,'只能选择一个'): build_forecast(self.d.history,self.d.assumptions)
 def test_workbooks_reopen(self):
  with tempfile.TemporaryDirectory() as d:
   i,r=generate_demo_artifacts(d); self.assertEqual(load_workbook(i).sheetnames,['历史数据','情景假设','使用说明']); self.assertEqual(load_workbook(r).sheetnames,REPORT_SHEETS)
