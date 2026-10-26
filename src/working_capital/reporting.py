from pathlib import Path
import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.chart import LineChart,BarChart,Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from .model import summarize_scenarios
from .evidence import profit_cash_bridge,one_factor_sensitivity

REPORT_SHEETS=['管理摘要','情景比较','现金流预测','营运资金','应收账款','库存与采购','应付账款','现金转换周期','资金缺口预警','历史数据','情景假设','数据质量','利润现金桥','单因素敏感性']
H={'scenario':'情景','month':'月份','revenue_mn':'收入（百万元）','collections_mn':'客户回款（百万元）','cogs_mn':'销售成本（百万元）','opening_inventory_mn':'期初库存（百万元）','purchases_mn':'采购（百万元）','ending_inventory_mn':'期末库存（百万元）','opening_ar_mn':'期初应收（百万元）','ending_ar_mn':'期末应收（百万元）','opening_ap_mn':'期初应付（百万元）','ending_ap_mn':'期末应付（百万元）','supplier_payments_mn':'供应商付款（百万元）','opex_mn':'经营费用（百万元）','operating_cashflow_mn':'经营现金流（百万元）','opening_cash_mn':'期初现金（百万元）','ending_cash_mn':'期末现金（百万元）','funding_gap_mn':'资金缺口（百万元）','dso':'DSO（天）','dio':'DIO（天）','dpo':'DPO（天）','ccc':'CCC（天）','minimum_cash_mn':'最低现金（百万元）','maximum_funding_gap_mn':'最大资金缺口（百万元）','average_ccc':'平均CCC（天）','revenue_growth':'收入增长率','cogs_rate':'销售成本率','opex_rate':'经营费用率'}
H.update({'target_inventory_mn':'目标库存（百万元）','inventory_above_target_mn':'超目标库存（百万元）'})
H['net_working_capital_mn']='净营运资金（百万元）'
H.update({'ar_mn':'期末应收（百万元）','inventory_mn':'期末库存（百万元）','ap_mn':'期末应付（百万元）','cash_mn':'期末现金（百万元）'})

def write(ws,df):
    for j,c in enumerate(df.columns,1):
        x=ws.cell(1,j,H.get(c,c)); x.font=Font(color='FFFFFF',bold=True); x.fill=PatternFill('solid',fgColor='24557A')
    for i,row in enumerate(df.itertuples(index=False,name=None),2):
        for j,v in enumerate(row,1): ws.cell(i,j,None if pd.isna(v) else v.to_pydatetime() if isinstance(v,pd.Timestamp) else v)
def fmt(ws):
    ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False; thin=Side(style='thin',color='D9E2F3')
    for row in ws.iter_rows():
        for c in row:
            c.border=Border(bottom=thin); c.alignment=Alignment(vertical='center',wrap_text=True)
            if hasattr(c.value,'year'): c.number_format='yyyy-mm'
            if isinstance(c.value,float): c.number_format='0.0%' if '率' in str(ws.cell(1,c.column).value) else '#,##0.00'
    for j in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(j)].width=min(max(12,max(len(str(ws.cell(i,j).value or '')) for i in range(1,ws.max_row+1))+2),25)

def generate_report(history,assumptions,path):
    r=summarize_scenarios(history,assumptions); f=r['forecast']; s=r['summary']; base=f[f.scenario=='基准'].copy()
    wb=Workbook(); wb.remove(wb.active); ws=wb.create_sheet('管理摘要'); ws.sheet_view.showGridLines=False
    ws['A1']='营运资金与现金流预测管理摘要'; ws['A1'].font=Font(size=20,bold=True,color='1F4E78')
    b=s.set_index('scenario'); metrics=[('历史期数',len(history)),('预测期数',12),('基准期末现金',b.loc['基准','ending_cash_mn']),('压力最低现金',b.loc['压力','minimum_cash_mn']),('压力最大资金缺口',b.loc['压力','maximum_funding_gap_mn']),('基准平均CCC',b.loc['基准','average_ccc'])]
    for i,(k,v) in enumerate(metrics,3): ws.cell(i,1,k).font=Font(bold=True); ws.cell(i,2,float(v))
    ws['A15']='零分母周转天数留空；多余库存逐月消耗，不假设退货、报废或负采购。'; ws['A16']='应收与应付目标受可用余额约束；CCC均值仅涵盖指标有定义的月份。'; ws['A11']='管理解读'; ws['A11'].font=Font(bold=True,color='1F4E78'); ws['A12']='收入增长可能同步增加应收与库存占用，利润增长不等于现金同步增长。'; ws['A13']='资金缺口必须单独披露，不能用自动融资掩盖。'; ws.column_dimensions['A'].width=66; ws.column_dimensions['B'].width=20
    cs=wb.create_sheet('情景比较'); write(cs,s); chart=BarChart(); chart.title='情景最低现金'; chart.add_data(Reference(cs,min_col=4,min_row=1,max_row=5),titles_from_data=True); chart.series[0].tx=SeriesLabel(strRef=StrRef(f="'{cs.title}'!$D$1")); chart.set_categories(Reference(cs,min_col=1,min_row=2,max_row=5)); cs.add_chart(chart,'I2')
    cash=wb.create_sheet('现金流预测'); write(cash,base[['month','opening_cash_mn','collections_mn','supplier_payments_mn','opex_mn','operating_cashflow_mn','ending_cash_mn','funding_gap_mn']]); line=LineChart(); line.title='基准现金余额'; line.add_data(Reference(cash,min_col=7,min_row=1,max_row=13),titles_from_data=True); line.series[0].tx=SeriesLabel(strRef=StrRef(f="'{cash.title}'!$G$1")); line.set_categories(Reference(cash,min_col=1,min_row=2,max_row=13)); cash.add_chart(line,'J2')
    wc=wb.create_sheet('营运资金'); x=base[['month','ending_ar_mn','ending_inventory_mn','ending_ap_mn']].copy(); x['net_working_capital_mn']=x.ending_ar_mn+x.ending_inventory_mn-x.ending_ap_mn; write(wc,x)
    ar=wb.create_sheet('应收账款'); write(ar,base[['month','revenue_mn','opening_ar_mn','collections_mn','ending_ar_mn','dso']])
    inv=wb.create_sheet('库存与采购'); write(inv,base[['month','cogs_mn','opening_inventory_mn','purchases_mn','ending_inventory_mn','target_inventory_mn','inventory_above_target_mn','dio']])
    ap=wb.create_sheet('应付账款'); write(ap,base[['month','purchases_mn','opening_ap_mn','supplier_payments_mn','ending_ap_mn','dpo']])
    ccc=wb.create_sheet('现金转换周期'); write(ccc,base[['month','dso','dio','dpo','ccc']])
    gap=wb.create_sheet('资金缺口预警'); g=f[f.funding_gap_mn>0][['scenario','month','ending_cash_mn','funding_gap_mn']]; write(gap,g if not g.empty else pd.DataFrame([{'scenario':'全部情景','month':pd.NaT,'ending_cash_mn':0.0,'funding_gap_mn':0.0}]))
    hist=wb.create_sheet('历史数据'); write(hist,history); ass=wb.create_sheet('情景假设'); write(ass,assumptions)
    q=wb.create_sheet('数据质量')
    residuals = {
        '应收滚动勾稽': f.opening_ar_mn+f.revenue_mn-f.collections_mn-f.ending_ar_mn,
        '库存滚动勾稽': f.opening_inventory_mn+f.purchases_mn-f.cogs_mn-f.ending_inventory_mn,
        '应付滚动勾稽': f.opening_ap_mn+f.purchases_mn-f.supplier_payments_mn-f.ending_ap_mn,
        '现金滚动勾稽': f.opening_cash_mn+f.operating_cashflow_mn-f.ending_cash_mn,
    }
    checks=[{'检查项':key,'结果':'通过' if value.abs().max()<1e-8 else '异常'} for key,value in residuals.items()]
    checks.append({'检查项':'CCC无定义月份（所有情景）','结果':str(int(f.ccc.isna().sum()))})
    write(q,pd.DataFrame(checks))
    bridge=profit_cash_bridge(f)
    write(wb.create_sheet('利润现金桥'),bridge)
    selected=assumptions[assumptions.scenario.eq('基准')].iloc[0].to_dict()
    write(wb.create_sheet('单因素敏感性'),one_factor_sensitivity(history,selected))
    ws['A18']='范围：合成经营现金情景；不含税款、资本开支和融资。金额单位：百万元。'
    ws['A19']='DPO按采购额，周转指标为月末余额/当月流量年化；不是实际回款或付款天数。'
    for sh in wb.worksheets: fmt(sh)
    out=Path(path); out.parent.mkdir(parents=True,exist_ok=True); wb.save(out); return out

def generate_demo_artifacts(folder):
    from .synthetic import build_synthetic_dataset
    d=build_synthetic_dataset(); p=Path(folder); p.mkdir(parents=True,exist_ok=True); inp=p/'合成营运资金输入.xlsx'
    with pd.ExcelWriter(inp,engine='openpyxl') as w:
        d.history.to_excel(w,sheet_name='历史数据',index=False); d.assumptions.to_excel(w,sheet_name='情景假设',index=False); pd.DataFrame({'说明':['全部数据由程序生成，不对应任何真实企业、客户、供应商或银行账户。']}).to_excel(w,sheet_name='使用说明',index=False)
    iw=load_workbook(inp)
    for ws in iw.worksheets:
        for cell in ws[1]:
            cell.value=H.get(str(cell.value),cell.value); cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='24557A')
        fmt(ws)
    iw.save(inp)
    return inp,generate_report(d.history,d.assumptions,p/'营运资金与现金流预测报告.xlsx')

# Labels used by the newly integrated evidence schedules.
H.update({'operating_profit_mn':'经营利润（百万元）','ar_cash_effect_mn':'应收变动现金影响（百万元）',
          'inventory_cash_effect_mn':'库存变动现金影响（百万元）','ap_cash_effect_mn':'应付变动现金影响（百万元）',
          'indirect_ocf_mn':'间接法经营现金（百万元）','direct_ocf_mn':'直接法经营现金（百万元）',
          'reconciliation_difference_mn':'现金勾稽差异（百万元）','driver':'单一变动参数','base_days':'基准天数',
          'scenario_days':'情景天数','delta_days':'变化天数','change_vs_base_mn':'期末现金较基准变化（百万元）',
          'maximum_gap_mn':'最大资金缺口（百万元）','zero_purchase_months':'零采购月份数',
          'above_target_inventory_months':'超目标库存月份数'})
